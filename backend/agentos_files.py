"""AgentOS File Center: locations (OneDrive, local folders), browsing, search, text extraction, AI tools."""
import collections, json, os, re, secrets, subprocess, time
from pathlib import Path

from fastapi import APIRouter, HTTPException
from pydantic import BaseModel

MODES = ("school", "dev")
TEXT_EXT = {".md", ".txt", ".json", ".csv", ".tsv", ".py", ".js", ".ts", ".jsx", ".tsx", ".html", ".htm", ".css", ".yaml", ".yml", ".toml", ".ini", ".log", ".xml", ".sh", ".java", ".c", ".cpp", ".h", ".go", ".rs", ".rb", ".swift", ".sql", ".tex", ".cfg", ".conf"}
OLD_OFFICE = {".doc", ".ppt", ".xls"}
BLOCK_NAMES = {".env", ".ssh", ".aws", ".gnupg", ".netrc", "id_rsa", "id_ed25519", "Keychains", "browser_profile", "Cookies", "Login Data"}
SKIP_DIRS = {"node_modules", ".git", ".venv", "venv", "__pycache__", "Library", ".Trash", "site-packages", "Caches"}
TCC = "macOS blocked access to this folder. Open System Settings > Privacy & Security > Full Disk Access (or Files & Folders), turn on your terminal app (Terminal), then restart AgentOS and try again."
MAX_UI = 60000


class LocBody(BaseModel):
    mode: str
    path: str
    name: str = ""


class RevealBody(BaseModel):
    mode: str
    loc: str
    path: str = ""


def isdir(p):
    try:
        return Path(p).is_dir()
    except OSError:
        return False


def is_cloud(path):
    return "CloudStorage" in str(path) or "/OneDrive" in str(path)


def _extract(p):
    ext = p.suffix.lower()
    try:
        size = p.stat().st_size
    except PermissionError:
        return "", TCC
    except OSError as e:
        return "", "Cannot access file: " + str(e)[:120]
    try:
        if ext in TEXT_EXT or ext == "":
            with open(p, "rb") as f:
                raw = f.read(2000000)
            if b"\x00" in raw[:4096]:
                return "", "Binary file, cannot show as text."
            return raw.decode("utf-8", errors="replace"), ("Showing the first 2 MB." if size > 2000000 else "")
        if ext in OLD_OFFICE:
            return "", "Old Office format (" + ext + "). Save it as docx, pptx or xlsx to read it."
        if size > 80000000:
            return "", "File too large to read (over 80 MB)."
        if ext == ".pdf":
            from pypdf import PdfReader
            r = PdfReader(str(p))
            out = []
            n = len(r.pages)
            for i, pg in enumerate(r.pages[:80]):
                t = pg.extract_text() or ""
                if t.strip():
                    out.append("--- page " + str(i + 1) + " ---\n" + t)
            if not out:
                return "", "No extractable text (this PDF may be scanned images)."
            return "\n\n".join(out), ("Read the first 80 of " + str(n) + " pages." if n > 80 else "")
        if ext == ".docx":
            from docx import Document
            d = Document(str(p))
            parts = [x.text for x in d.paragraphs if x.text.strip()]
            for t in d.tables:
                for row in t.rows:
                    parts.append(" | ".join(c.text.strip() for c in row.cells))
            return "\n".join(parts), ""
        if ext == ".pptx":
            from pptx import Presentation
            pr = Presentation(str(p))
            out = []
            for i, s in enumerate(pr.slides, 1):
                t = []
                for sh in s.shapes:
                    if getattr(sh, "has_text_frame", False) and sh.has_text_frame:
                        tx = sh.text_frame.text.strip()
                        if tx:
                            t.append(tx)
                out.append("--- slide " + str(i) + " ---\n" + "\n".join(t))
            return "\n\n".join(out), ""
        if ext == ".xlsx":
            from openpyxl import load_workbook
            wb = load_workbook(str(p), read_only=True, data_only=True)
            out = []
            for ws in wb.worksheets[:6]:
                out.append("--- sheet: " + ws.title + " ---")
                for ri, row in enumerate(ws.iter_rows(values_only=True)):
                    if ri >= 300:
                        out.append("...")
                        break
                    out.append("\t".join("" if v is None else str(v) for v in row))
            wb.close()
            return "\n".join(out), ""
    except PermissionError:
        return "", TCC
    except ImportError:
        return "", "A reader library is missing. In Terminal run: ~/Desktop/AgentOS/.venv/bin/pip install pypdf python-docx python-pptx openpyxl"
    except Exception as e:
        return "", "Could not read this file: " + str(e)[:150]
    return "", "Unsupported file type (" + (ext or "none") + ")."


def extract(p):
    text, note = _extract(Path(p))
    return text[:500000], note


def setup(app, G):
    ROOT, DATA = G["ROOT"], G["DATA"]
    LOCF = DATA / "locations.json"
    BRAIN = ROOT / "brain"
    rt = APIRouter()

    def M(m):
        if m not in MODES:
            raise HTTPException(400, "mode must be school or dev")
        return m

    def load():
        try:
            return json.loads(LOCF.read_text(encoding="utf-8"))
        except Exception:
            return {}

    def locs(mode):
        return load().get(M(mode), [])

    def find_loc(mode, lid):
        for l in locs(mode):
            if l["id"] == lid:
                return l
        raise HTTPException(404, "Unknown location")

    def inside(base, p):
        return p == base or base in p.parents

    def blocked(p):
        p = Path(p)
        if any(x in BLOCK_NAMES for x in p.parts):
            return True
        if inside(ROOT, p):
            return not inside(BRAIN, p)
        return False

    def resolve_in(l, rel):
        base = Path(l["path"]).resolve()
        p = (base / (rel or "").lstrip("/")).resolve()
        if not inside(base, p):
            raise HTTPException(403, "Outside this location")
        if blocked(p):
            raise HTTPException(403, "Protected path")
        return base, p

    def suggestions():
        home = Path.home()
        out = []
        cs = home / "Library" / "CloudStorage"
        try:
            for p in sorted(cs.glob("OneDrive*")):
                if isdir(p):
                    out.append({"name": ("OneDrive: " + p.name[9:]) if p.name.startswith("OneDrive-") else "OneDrive", "path": str(p.resolve())})
        except OSError:
            pass
        try:
            for p in sorted(home.glob("OneDrive*")):
                if isdir(p) and str(p.resolve()) not in [o["path"] for o in out]:
                    out.append({"name": p.name, "path": str(p.resolve())})
        except OSError:
            pass
        for n in ("Desktop", "Documents", "Downloads"):
            if isdir(home / n):
                out.append({"name": n, "path": str((home / n).resolve())})
        out.append({"name": "Home folder", "path": str(home.resolve())})
        try:
            for p in sorted(Path("/Volumes").iterdir())[:6]:
                if isdir(p) and str(p.resolve()) != "/":
                    out.append({"name": "Drive: " + p.name, "path": str(p.resolve())})
        except OSError:
            pass
        return out

    # ---------- search / tools ----------
    def do_search(base, start_rel, q, contents, limit, budget):
        base = Path(base).resolve()
        start = (base / (start_rel or "").lstrip("/")).resolve()
        if not inside(base, start):
            return [], False
        words = [w for w in q.lower().split() if w]
        phrase = q.lower().strip()
        t0, res, trunc = time.time(), [], False
        dq = collections.deque([(start, 0)])
        while dq:
            d, depth = dq.popleft()
            try:
                entries = sorted(os.scandir(d), key=lambda x: x.name.lower())
            except PermissionError:
                if depth == 0:
                    raise
                continue
            except OSError:
                continue
            for e in entries:
                if time.time() - t0 > budget or len(res) >= limit:
                    trunc = True
                    dq.clear()
                    break
                name = e.name
                if name.startswith("."):
                    continue
                try:
                    isd = e.is_dir(follow_symlinks=False)
                except OSError:
                    continue
                ep = Path(e.path)
                if blocked(ep):
                    continue
                rel = ep.relative_to(base).as_posix()
                hit = bool(words) and all(w in name.lower() for w in words)
                if isd:
                    if name in SKIP_DIRS:
                        continue
                    if depth < 7:
                        dq.append((ep, depth + 1))
                    if hit:
                        res.append({"path": rel, "name": name, "is_dir": True, "hit": ""})
                    continue
                if hit:
                    res.append({"path": rel, "name": name, "is_dir": False, "hit": ""})
                elif contents and ep.suffix.lower() in TEXT_EXT:
                    try:
                        if e.stat().st_size > 1000000:
                            continue
                        txt = ep.read_text(encoding="utf-8", errors="replace")
                    except OSError:
                        continue
                    low = txt.lower()
                    i = low.find(phrase) if phrase else -1
                    if i >= 0:
                        a = txt.rfind("\n", 0, i) + 1
                        b = txt.find("\n", i)
                        res.append({"path": rel, "name": name, "is_dir": False, "hit": txt[a:(b if b > 0 else len(txt))].strip()[:160]})
        return res, trunc

    def fs_roots(mode):
        out = []
        for l in locs(mode):
            if isdir(l["path"]):
                out.append(Path(l["path"]).resolve())
        return out

    def fs_tool(mode, name, args):
        if name == "list_locations":
            lines = ["Notes folder (read and write): use relative paths such as Biology/Cells.md"]
            for l in locs(mode):
                lines.append("- " + l["name"] + " (read-only): " + l["path"])
            if len(lines) == 1:
                lines.append("No other locations are added yet. The user can add OneDrive, Documents and more in the Second Brain + Files tab.")
            return "\n".join(lines)
        if name == "find_files":
            q = str(args.get("query", "")).strip()
            if not q:
                return "Empty query."
            want = str(args.get("location", "")).strip().lower()
            cands = [l for l in locs(mode) if not want or want in l["name"].lower()]
            if not cands:
                return "No matching location. Call list_locations first."
            contents = str(args.get("contents", "")).lower() in ("true", "1", "yes")
            out = []
            for l in cands:
                base = Path(l["path"])
                if not isdir(base):
                    continue
                try:
                    res, trunc = do_search(base, "", q, contents and not is_cloud(l["path"]), 15, 5.0)
                except PermissionError:
                    out.append(l["name"] + ": " + TCC)
                    continue
                root = base.resolve()
                for r in res:
                    out.append(str(root / r["path"]) + ("  (folder)" if r["is_dir"] else "") + (("  -> " + r["hit"]) if r["hit"] else ""))
            return "\n".join(out[:40]) or "No matches."
        return None

    # ---------- API ----------
    @rt.get("/api/fs/locations")
    def fs_locations(mode: str):
        cur = locs(mode)
        have = {l["path"] for l in cur}
        return {"locations": [dict(l, exists=isdir(l["path"])) for l in cur], "suggestions": [s for s in suggestions() if s["path"] not in have]}

    @rt.post("/api/fs/locations")
    def fs_add(b: LocBody):
        mode = M(b.mode)
        p = Path(os.path.expanduser(b.path.strip())).resolve()
        if not isdir(p):
            raise HTTPException(404, "That folder does not exist: " + str(p))
        if blocked(p):
            raise HTTPException(403, "That folder is protected and cannot be added.")
        try:
            os.listdir(p)
        except PermissionError:
            raise HTTPException(403, TCC)
        except OSError as e:
            raise HTTPException(400, str(e))
        d = load()
        lst = d.get(mode, [])
        for l in lst:
            if l["path"] == str(p):
                return {"location": l}
        loc = {"id": secrets.token_hex(3), "name": (b.name.strip() or p.name or "Folder")[:40], "path": str(p)}
        lst.append(loc)
        d[mode] = lst
        LOCF.write_text(json.dumps(d, indent=1), encoding="utf-8")
        return {"location": loc}

    @rt.delete("/api/fs/locations")
    def fs_remove(mode: str, id: str):
        d = load()
        d[M(mode)] = [l for l in d.get(mode, []) if l["id"] != id]
        LOCF.write_text(json.dumps(d, indent=1), encoding="utf-8")
        return {"removed": id}

    @rt.get("/api/fs/list")
    def fs_list(mode: str, loc: str, path: str = ""):
        l = find_loc(mode, loc)
        base, p = resolve_in(l, path)
        if not p.is_dir():
            raise HTTPException(404, "Not a folder")
        try:
            entries = list(os.scandir(p))
        except PermissionError:
            raise HTTPException(403, TCC)
        except OSError as e:
            raise HTTPException(500, str(e))
        items = []
        for e in entries:
            if e.name.startswith("."):
                continue
            try:
                d = e.is_dir()
                st = e.stat()
            except OSError:
                continue
            if blocked(Path(e.path)):
                continue
            items.append({"name": e.name, "is_dir": d, "size": None if d else st.st_size, "mtime": st.st_mtime})
        items.sort(key=lambda x: (not x["is_dir"], x["name"].lower()))
        return {"items": items[:1500]}

    @rt.get("/api/fs/read")
    def fs_read(mode: str, loc: str, path: str):
        l = find_loc(mode, loc)
        base, p = resolve_in(l, path)
        if not p.is_file():
            raise HTTPException(404, "Not a file")
        text, note = extract(p)
        return {"name": p.name, "abs": str(p), "size": p.stat().st_size, "text": text[:MAX_UI], "truncated": len(text) > MAX_UI, "note": note}

    @rt.get("/api/fs/search")
    def fs_search(mode: str, loc: str, q: str, path: str = ""):
        l = find_loc(mode, loc)
        base, p = resolve_in(l, path)
        if not q.strip():
            return {"results": [], "truncated": False}
        try:
            res, trunc = do_search(base, p.relative_to(base).as_posix() if p != base else "", q, not is_cloud(l["path"]), 60, 8.0)
        except PermissionError:
            raise HTTPException(403, TCC)
        return {"results": res, "truncated": trunc}

    @rt.post("/api/fs/reveal")
    def fs_reveal(b: RevealBody):
        l = find_loc(b.mode, b.loc)
        base, p = resolve_in(l, b.path)
        subprocess.run(["open", "-R", str(p)], check=False)
        return {"ok": True}

    # ---------- AI tools ----------
    import agentos_ext
    NEW = [
        {"type": "function", "function": {"name": "read_file", "description": "Read a file as text. Works on text, markdown, code, PDF, Word (docx), PowerPoint (pptx) and Excel (xlsx). A relative path is inside the notes folder; use the full path for files in File Center locations. For long files, call again with start set to continue.",
            "parameters": {"type": "object", "properties": {"path": {"type": "string"}, "start": {"type": "integer", "description": "character offset to continue reading, default 0"}}, "required": ["path"]}}},
        {"type": "function", "function": {"name": "list_locations", "description": "List the folders the user has given you read access to (OneDrive, Documents, Desktop and so on) with their full paths.",
            "parameters": {"type": "object", "properties": {}}}},
        {"type": "function", "function": {"name": "find_files", "description": "Search file names (and optionally the text inside text files) across the user's File Center locations. Returns full paths you can pass to read_file.",
            "parameters": {"type": "object", "properties": {"query": {"type": "string"}, "location": {"type": "string", "description": "optional location name, such as OneDrive"}, "contents": {"type": "boolean", "description": "also search inside text files (not available for OneDrive)"}}, "required": ["query"]}}},
    ]
    drop = ("read_file", "list_locations", "find_files")
    agentos_ext.TOOLS[:] = [t for t in agentos_ext.TOOLS if t["function"]["name"] not in drop] + NEW

    G["fs_roots"] = fs_roots
    G["fs_blocked"] = blocked
    G["fs_extract"] = extract
    G["fs_tool"] = fs_tool
    app.include_router(rt)
